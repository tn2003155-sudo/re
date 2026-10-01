"""レシート画像を読み取り、日付・店名・購入品目をExcelにまとめるツール.

使い方:
    python receipt_to_xlsx.py receipts/ -o receipts.xlsx
    python receipt_to_xlsx.py img1.jpg img2.png -o out.xlsx
    python receipt_to_xlsx.py receipts/ --names names.csv   # 代表品名リストに寄せる
"""

from __future__ import annotations

import argparse
import base64
import copy
import csv
import io
import json
import sys
from pathlib import Path

import anthropic
from openpyxl import Workbook
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter
from PIL import Image, ImageOps

MODEL = "claude-opus-5-5"
IMAGE_EXTS = {".jpg", ".jpeg", ".png", ".gif", ".webp"}
MAX_IMAGE_BYTES = 4_500_000  # API上限 5MB に余裕を持たせる
MAX_LONG_EDGE = 2400

RECEIPT_SCHEMA = {
    "type": "object",
    "properties": {
        "date": {
            "type": "string",
            "description": "購入日 (YYYY-MM-DD)。読み取れない場合は空文字",
        },
        "store_name": {"type": "string", "description": "店名。読み取れない場合は空文字"},
        "items": {
            "type": "array",
            "items": {
                "type": "object",
                "properties": {
                    "name": {"type": "string", "description": "シンプルな代表品名"},
                    "original_name": {"type": "string", "description": "レシート上の表記そのまま"},
                    "quantity": {"type": "number", "description": "数量 (不明なら1)"},
                    "unit_price": {"type": "number", "description": "単価 (円)"},
                    "amount": {"type": "number", "description": "金額 (円、値引き後)"},
                },
                "required": ["name", "original_name", "quantity", "unit_price", "amount"],
                "additionalProperties": False,
            },
        },
        "total": {"type": "number", "description": "合計金額 (円、税込)"},
    },
    "required": ["date", "store_name", "items", "total"],
    "additionalProperties": False,
}

PROMPT = """このレシート画像から、購入日・店名・購入品目・合計金額を抽出してください。
- 日付は西暦 YYYY-MM-DD 形式に変換する(和暦や「26/10/01」のような表記も変換)。
- 品目には商品のみを含め、小計・合計・税・お預り・お釣り・ポイントなどの行は含めない。
- 値引き行は直前の商品の金額に反映する。
- original_name にはレシートの表記どおりに書き起こす。
- name には、ブランド名・容量・産地・略語などを省いたシンプルな代表品名を書く
  (例: 「明治おいしい牛乳1000ml」→「牛乳」、「国産若鶏もも肉」→「鶏肉」)。"""

OTHER = "その他"
HEADER_WORDS = {"品名", "品目", "代表品名", "購入品名", "商品名"}


def load_names(path: Path) -> list[str]:
    """代表品名リストを読み込む (CSV/TXT の1列目、空行・重複は除く)."""
    names: list[str] = []
    with path.open(encoding="utf-8-sig", newline="") as f:
        for row in csv.reader(f):
            name = row[0].strip() if row else ""
            if name and name not in names and name not in HEADER_WORDS:
                names.append(name)
    return names


def build_request(names: list[str] | None) -> tuple[dict, str]:
    """品名リストがあれば name を列挙型に制限したスキーマとプロンプトを返す."""
    if not names:
        return RECEIPT_SCHEMA, PROMPT
    schema = copy.deepcopy(RECEIPT_SCHEMA)
    choices = names + ([OTHER] if OTHER not in names else [])
    schema["properties"]["items"]["items"]["properties"]["name"]["enum"] = choices
    prompt = PROMPT + f"""
- name は次の代表品名リストから最も近いものを1つ選ぶ。どれにも当てはまらない場合は「{OTHER}」とする。
代表品名リスト: {"、".join(names)}"""
    return schema, prompt


def load_image(path: Path) -> tuple[str, str]:
    """画像を(必要なら縮小して)base64化し、(media_type, data) を返す."""
    raw = path.read_bytes()
    suffix = path.suffix.lower()
    media_type = {
        ".jpg": "image/jpeg",
        ".jpeg": "image/jpeg",
        ".png": "image/png",
        ".gif": "image/gif",
        ".webp": "image/webp",
    }[suffix]

    img = Image.open(io.BytesIO(raw))
    if len(raw) > MAX_IMAGE_BYTES or max(img.size) > MAX_LONG_EDGE:
        img = ImageOps.exif_transpose(img).convert("RGB")
        img.thumbnail((MAX_LONG_EDGE, MAX_LONG_EDGE))
        quality = 90
        while True:
            buf = io.BytesIO()
            img.save(buf, format="JPEG", quality=quality)
            if buf.tell() <= MAX_IMAGE_BYTES or quality <= 50:
                break
            quality -= 10
        raw, media_type = buf.getvalue(), "image/jpeg"

    return media_type, base64.standard_b64encode(raw).decode("utf-8")


def extract_receipt(client: anthropic.Anthropic, path: Path, schema: dict, prompt: str) -> dict:
    media_type, data = load_image(path)
    response = client.beta.messages.create(
        model=MODEL,
        max_tokens=16000,
        betas=["server-side-fallback-2026-07-01"],
        fallbacks="default",
        output_config={
            "effort": "medium",
            "format": {"type": "json_schema", "schema": schema},
        },
        messages=[{
            "role": "user",
            "content": [
                {"type": "image", "source": {"type": "base64", "media_type": media_type, "data": data}},
                {"type": "text", "text": prompt},
            ],
        }],
    )
    if response.stop_reason == "refusal":
        raise RuntimeError("モデルが応答を拒否しました")
    if response.stop_reason == "max_tokens":
        raise RuntimeError("出力が途中で切れました")
    text = next(b.text for b in response.content if b.type == "text")
    return json.loads(text)


def collect_images(inputs: list[str]) -> list[Path]:
    paths: list[Path] = []
    for arg in inputs:
        p = Path(arg)
        if p.is_dir():
            paths.extend(sorted(f for f in p.iterdir() if f.suffix.lower() in IMAGE_EXTS))
        elif p.suffix.lower() in IMAGE_EXTS:
            paths.append(p)
        else:
            print(f"スキップ (対応していない形式): {p}", file=sys.stderr)
    return paths


def style_header(ws, widths: list[int]) -> None:
    for i, width in enumerate(widths, start=1):
        cell = ws.cell(row=1, column=i)
        cell.font = Font(bold=True, color="FFFFFF")
        cell.fill = PatternFill("solid", fgColor="4472C4")
        cell.alignment = Alignment(horizontal="center")
        ws.column_dimensions[get_column_letter(i)].width = width
    ws.freeze_panes = "A2"
    ws.auto_filter.ref = ws.dimensions


def write_workbook(results: list[tuple[Path, dict]], out: Path) -> None:
    results = sorted(results, key=lambda r: (r[1]["date"] or "9999", r[1]["store_name"]))
    wb = Workbook()

    ws = wb.active
    ws.title = "明細"
    ws.append(["日付", "店名", "品名", "数量", "単価", "金額", "レシート表記", "ファイル"])
    for path, r in results:
        for item in r["items"]:
            ws.append([r["date"], r["store_name"], item["name"], item["quantity"],
                       item["unit_price"], item["amount"], item["original_name"], path.name])
    for row in ws.iter_rows(min_row=2, min_col=5, max_col=6):
        for cell in row:
            cell.number_format = "#,##0"
    style_header(ws, [12, 24, 16, 8, 10, 10, 36, 24])

    summary = wb.create_sheet("レシート一覧")
    summary.append(["日付", "店名", "品目数", "品目合計", "レシート合計", "差額", "ファイル"])
    for i, (path, r) in enumerate(results, start=2):
        summary.append([r["date"], r["store_name"], len(r["items"]),
                        sum(item["amount"] for item in r["items"]), r["total"],
                        f"=E{i}-D{i}", path.name])
    for row in summary.iter_rows(min_row=2, min_col=4, max_col=6):
        for cell in row:
            cell.number_format = "#,##0"
    style_header(summary, [12, 24, 8, 12, 12, 10, 24])

    wb.save(out)


def main() -> int:
    parser = argparse.ArgumentParser(description="レシート画像をExcelにまとめます")
    parser.add_argument("inputs", nargs="+", help="画像ファイルまたはフォルダ")
    parser.add_argument("-o", "--output", default="receipts.xlsx", help="出力ファイル名")
    parser.add_argument("--names", type=Path,
                        help="代表品名リスト (CSV/TXT、1列目を使用)。品名をこのリストに寄せる")
    args = parser.parse_args()

    names = load_names(args.names) if args.names else None
    schema, prompt = build_request(names)

    images = collect_images(args.inputs)
    if not images:
        print("画像が見つかりませんでした", file=sys.stderr)
        return 1

    client = anthropic.Anthropic()
    results: list[tuple[Path, dict]] = []
    for path in images:
        print(f"読み取り中: {path}")
        try:
            results.append((path, extract_receipt(client, path, schema, prompt)))
        except (anthropic.APIConnectionError, anthropic.APIStatusError, RuntimeError) as e:
            print(f"  失敗: {e}", file=sys.stderr)

    if not results:
        return 1
    write_workbook(results, Path(args.output))
    print(f"{len(results)}/{len(images)} 件を {args.output} に保存しました")
    return 0


if __name__ == "__main__":
    sys.exit(main())
